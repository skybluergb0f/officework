import hashlib
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from server.workbook import inspect_workbook, workbook_path_for_date, write_workbook_record


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA = PROJECT_ROOT / "data"


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class DataWorkbookTests(unittest.TestCase):
    def test_resolves_only_the_supplied_data_month_copy(self):
        path = workbook_path_for_date("2026-10-01", DATA)
        self.assertEqual(path, DATA / "2026년 10월 수변전 일지 - 본관.xlsx")
        self.assertTrue(path.is_file())

    def test_data_month_copies_have_daily_sheets_and_print_parts(self):
        for month in ("09", "10"):
            with self.subTest(month=month):
                path = workbook_path_for_date(f"2026-{month}-01", DATA)
                before = sha256(path)
                report = inspect_workbook(path)
                self.assertIn("수식(월변경)", report["sheet_names"])
                self.assertIn("01일", report["sheet_names"])
                self.assertEqual(report["day_sheet_count"], 31)
                self.assertGreaterEqual(report["printer_settings_count"], 31)
                self.assertEqual(sha256(path), before, "inspection must not modify the data copy")

    def test_day_must_be_a_real_calendar_date(self):
        with self.assertRaises(ValueError):
            workbook_path_for_date("2026-02-30", DATA)

    def test_projection_on_a_copy_changes_only_target_sheet_and_keeps_formula(self):
        source = workbook_path_for_date("2026-09-02", DATA)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            target = root / source.name
            backup_dir = root / "backups"
            shutil.copy2(source, target)
            with zipfile.ZipFile(target) as archive:
                before = {name: archive.read(name) for name in archive.namelist()}
                settings_before = {name: data for name, data in before.items() if name.startswith("xl/printerSettings/")}
            original_hash = sha256(target)
            cached_book = load_workbook(target, data_only=True)
            formula_book = load_workbook(target, data_only=False)
            try:
                cached_previous = cached_book["02일"]["N39"].value
                formula = formula_book["02일"]["N39"].value
            finally:
                cached_book.close()
                formula_book.close()
            self.assertIsInstance(formula, str)
            self.assertTrue(formula.startswith("="))

            mismatch = cached_previous + 1 if isinstance(cached_previous, (int, float)) else "다른 값"
            with tempfile.TemporaryDirectory() as mismatch_tmp:
                mismatch_target = Path(mismatch_tmp) / source.name
                mismatch_backups = Path(mismatch_tmp) / "backups"
                shutil.copy2(source, mismatch_target)
                mismatch_hash = sha256(mismatch_target)
                with self.assertRaisesRegex(ValueError, "formula-protected cell value mismatch"):
                    write_workbook_record(mismatch_target, "2026-09-02", {
                        "date": "2026-09-02",
                        "operator": "테스터",
                        "observations": {"08:00": {"main": {"kv": 22.9}}},
                        "meters": {"previous": [mismatch]},
                    }, mismatch_backups)
                self.assertEqual(sha256(mismatch_target), mismatch_hash)
                self.assertFalse(mismatch_backups.exists())

            result = write_workbook_record(target, "2026-09-02", {
                "date": "2026-09-02",
                "operator": "테스터",
                "observations": {"08:00": {"main": {"kv": 22.9}}},
                "meters": {"previous": [cached_previous]},
            }, backup_dir)

            with zipfile.ZipFile(target) as archive:
                after = {name: archive.read(name) for name in archive.namelist()}
            changed = {name for name in before if before[name] != after[name]}
            self.assertEqual(changed, {result["worksheet_part"], result["summary_part"]})
            self.assertEqual(settings_before, {name: data for name, data in after.items() if name.startswith("xl/printerSettings/")})
            self.assertEqual(sha256(result["backup"]), original_hash)
            verify = load_workbook(target, data_only=True)
            formulas_after = load_workbook(target, data_only=False)
            try:
                self.assertEqual(verify["02일"]["F10"].value, 22.9)
                self.assertEqual(verify["02일"]["N39"].value, cached_previous)
                self.assertEqual(verify["수식(월변경)"]["J3"].value, "테스터")
                self.assertEqual(formulas_after["02일"]["N39"].value, formula)
            finally:
                verify.close()
                formulas_after.close()
            self.assertEqual(result["preserved_formula_cells"], ["N39"])
            synced_hash = sha256(target)
            backup_count = len(list(backup_dir.glob("*.xlsx")))
            duplicate = write_workbook_record(target, "2026-09-02", {
                "date": "2026-09-02",
                "operator": "테스터",
                "observations": {"08:00": {"main": {"kv": 22.9}}},
                "meters": {"previous": [cached_previous]},
            }, backup_dir)
            self.assertTrue(duplicate["unchanged"])
            self.assertIsNone(duplicate["backup"])
            self.assertEqual(sha256(target), synced_hash)
            self.assertEqual(len(list(backup_dir.glob("*.xlsx"))), backup_count)


if __name__ == "__main__":
    unittest.main()
