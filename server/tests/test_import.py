import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from server.db import connect, initialize_database, load_record, save_record
from server.workbook import import_workbook_month


PROJECT_ROOT = Path(__file__).resolve().parents[2]
SEPTEMBER_COPY = PROJECT_ROOT / "data" / "2026년 9월 수변전 일지 - 본관.xlsx"


def sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


class SeptemberImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_path = Path(self.temp.name) / "import-test.sqlite3"
        initialize_database(self.db_path)
        self.db = connect(self.db_path)

    def tearDown(self):
        self.db.close()
        self.temp.cleanup()

    def test_reads_september_data_copy_maps_fields_and_persists_all_days(self):
        before_hash = sha256(SEPTEMBER_COPY)
        report = import_workbook_month(self.db, SEPTEMBER_COPY.parent, "2026-09")
        self.assertEqual(report["imported_days"], 30)
        self.assertEqual(report["source"], str(SEPTEMBER_COPY.resolve()))
        self.assertEqual(report["warning_count"], 1)
        self.assertEqual(report["warnings"][0], {"date": "2026-09-14", "time": "14:00", "group": "transformer", "field": "hvac", "source_text": "45..9"})

        source = load_workbook(SEPTEMBER_COPY, read_only=False, data_only=True)
        try:
            source_day = source["01일"]
            source_summary = source["수식(월변경)"]
            expected = {
                ("08:00", "main", "kv"): source_day["F10"].value,
                ("08:00", "vcb", "kv"): source_day["AH10"].value,
                ("08:00", "chiller", "chillerKw"): source_day["R20"].value,
                ("08:00", "transformer", "lighting"): source_day["AQ20"].value,
                ("08:00", "secondary", "lowLightingV"): source_day["F30"].value,
                ("08:00", "secondary", "rectifierV"): source_day["BP20"].value,
                ("08:00", "secondary", "batteryA"): source_day["CC20"].value,
            }
            meter_current = source_day["N38"].value
            substation_close = source_day["BK28"].value
            industrial_close = source_day["CB28"].value
            operator = source_summary["J2"].value
        finally:
            source.close()
        saved = load_record(self.db, "2026-09-01")
        self.assertNotIn("trTemp", saved["observations"]["08:00"]["transformer"])
        for (time, group, field), value in expected.items():
            with self.subTest(time=time, group=group, field=field):
                self.assertEqual(saved["observations"][time][group][field], value)
        self.assertEqual(saved["meters"]["current"][0], meter_current)
        self.assertEqual(saved["meters"]["monthlyClose"]["substation"]["current"][0], substation_close)
        self.assertEqual(saved["meters"]["monthlyClose"]["industrial"]["previous"][0], industrial_close)
        self.assertEqual(saved["operator"], operator)
        self.assertEqual(saved["meters"]["current"][5:], [None, None])
        cached_source = load_workbook(SEPTEMBER_COPY, read_only=False, data_only=True)
        try:
            formula_cache = cached_source["02일"]["N39"].value
        finally:
            cached_source.close()
        second_day = load_record(self.db, "2026-09-02")
        self.assertEqual(second_day["meters"]["previous"][0], formula_cache)
        invalid_record = load_record(self.db, "2026-09-14")
        self.assertEqual(invalid_record["observations"]["14:00"]["transformer"]["hvac"], "45..9")
        stored_raw = self.db.execute("SELECT value,source_text FROM inspection_values WHERE record_date=? AND time_slot=? AND group_key=? AND field_key=?", ("2026-09-14", "14:00", "transformer", "hvac")).fetchone()
        self.assertIsNone(stored_raw["value"])
        self.assertEqual(stored_raw["source_text"], "45..9")
        self.assertEqual(sha256(SEPTEMBER_COPY), before_hash)
        self.assertEqual(self.db.execute("SELECT count(*) FROM daily_records").fetchone()[0], 30)

    def test_refuses_to_overwrite_existing_day_without_explicit_authorization(self):
        save_record(self.db, {"date": "2026-09-01", "operator": "Existing"})
        with self.assertRaisesRegex(ValueError, "already exist"):
            import_workbook_month(self.db, SEPTEMBER_COPY.parent, "2026-09")
        self.assertEqual(load_record(self.db, "2026-09-01")["operator"], "Existing")
        self.assertEqual(self.db.execute("SELECT count(*) FROM daily_records").fetchone()[0], 1)


if __name__ == "__main__":
    unittest.main()
